"""Convert the evaluation strong ML resume into a .docx the UI can upload."""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "evaluation" / "datasets" / "resumes" / "ml_strong.txt"
OUTPUT = ROOT / "sample_resume.docx"


def main() -> None:
    text = SOURCE.read_text(encoding="utf-8").strip()
    document = Document()
    for block in text.split("\n\n"):
        lines = [line for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        document.add_heading(lines[0], level=1)
        for line in lines[1:]:
            document.add_paragraph(line)
    document.save(OUTPUT)
    print(f"wrote {OUTPUT.name} from {SOURCE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
